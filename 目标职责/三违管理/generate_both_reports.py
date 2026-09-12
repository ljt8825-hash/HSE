import pandas as pd
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# 需要排除的人员名单
exclude_names = ['杨清华', '张和祥', '孙建栋', '谢俊章', '张万恒', '靳洪刚', '郭志强', '张卫新', '王臣祥', '李振恒', '张华平']

# 读取乡镇人员名单
personnel_df = pd.read_excel('乡镇人员名单（微信）.xlsx')

# 构建映射关系（只包含管理员，排除乡镇主管和指定人员）
name_to_wechat = {}
wechat_to_name = {}
township_to_people = {}
all_personnel = []

for idx, row in personnel_df.iterrows():
    duty = str(row.iloc[2]).strip()
    name = str(row.iloc[3]).strip()
    wechat = str(row.iloc[4]).strip()
    township = str(row.iloc[1]).strip()
    
    # 跳过乡镇主管
    if '乡镇主管' in duty:
        continue
    
    # 跳过指定要排除的人员
    if name in exclude_names:
        continue
    
    if name and name != 'nan':
        if wechat and wechat != 'nan':
            name_to_wechat[name] = wechat
            wechat_to_name[wechat] = name
        
        if township and township != 'nan':
            if township not in township_to_people:
                township_to_people[township] = []
            township_to_people[township].append(name)
            all_personnel.append({
                'township': township,
                'name': name,
                'wechat': wechat
            })

# 读取拍照巡检数据
inspection_df = pd.read_excel('拍照巡检 AI 识别.xlsx', header=None)

# 获取当前日期
now = datetime.now()
today_date_str = now.strftime('%Y/%m/%d')
today_month_day = f'{now.month}/{now.day}'

# 查找所有创建人（用于完成情况统计）
completed_map = {}  # {name: [dates]}
sanwei_map = {}  # {name: [{date, details}]}

for idx in range(2, len(inspection_df)):
    creator = str(inspection_df.iloc[idx, 7]).strip()
    inspection_date = str(inspection_df.iloc[idx, 3]).strip()
    tuoxie = str(inspection_df.iloc[idx, 4]).strip()
    gongzhuang = str(inspection_df.iloc[idx, 5]).strip()
    toutou = str(inspection_df.iloc[idx, 6]).strip()
    
    if creator and creator != 'nan':
        # 获取真实姓名
        name = creator
        if creator in wechat_to_name:
            name = wechat_to_name[creator]
        
        # 跳过被排除的人员
        if name in exclude_names:
            continue
        
        # 记录完成日期
        if name not in completed_map:
            completed_map[name] = set()
        
        # 转换日期格式
        try:
            if '/' in inspection_date:
                date_obj = datetime.strptime(inspection_date, '%Y/%m/%d')
                month_day = f'{date_obj.month}/{date_obj.day}'
                full_date = inspection_date
            elif '年' in inspection_date and '月' in inspection_date and '日' in inspection_date:
                year_part = inspection_date.split('年')[0]
                month_part = inspection_date.split('年')[1].split('月')[0]
                day_part = inspection_date.split('月')[1].split('日')[0]
                month_day = f'{int(month_part)}/{int(day_part)}'
                full_date = f'{year_part}/{int(month_part):02d}/{int(day_part):02d}'
            else:
                date_obj = datetime.strptime(inspection_date, '%Y-%m-%d')
                month_day = f'{date_obj.month}/{date_obj.day}'
                full_date = inspection_date
            
            completed_map[name].add(month_day)
        except:
            month_day = ''
            full_date = inspection_date
        
        # 检查是否有三违行为（只统计当天的）
        if full_date == today_date_str or month_day == today_month_day:
            has_sanwei = False
            sanwei_details = []
            
            if tuoxie == '是':
                has_sanwei = True
                sanwei_details.append('穿拖鞋')
            
            if gongzhuang == '是':
                has_sanwei = True
                sanwei_details.append('没穿工装')
            
            if toutou == '是':
                has_sanwei = True
                sanwei_details.append('没戴头盔')
            
            if has_sanwei:
                if name not in sanwei_map:
                    sanwei_map[name] = []
                sanwei_map[name].append({
                    'date': full_date,
                    'details': '、'.join(sanwei_details)
                })

# 获取当前日期和时间
today_date = now.strftime('%Y年%m月%d日')
current_time = now.strftime('%Y年%m月%d日 %H:%M')

# ==================== 生成文字性叙述报告 ====================
report_lines = []
report_lines.append('📊 【宁津项目"三违"管理】')
report_lines.append(f'📅 统计周期: {today_date}')
report_lines.append(f'🕒 生成时间: {current_time}')

# 已完成情况
report_lines.append('\n✅ 已完成巡检的乡镇：')
completed_townships = {}
for township, people in sorted(township_to_people.items()):
    completed_people = [p for p in people if p in completed_map]
    if completed_people:
        completed_townships[township] = completed_people

if completed_townships:
    for township, people in sorted(completed_townships.items()):
        people_str = '、'.join(people)
        report_lines.append(f'{township}：{people_str}。')
else:
    report_lines.append('无')

# 未完成情况
report_lines.append('\n❌ 未完成巡检的乡镇：')
incomplete_townships = {}
for township, people in sorted(township_to_people.items()):
    incomplete_people = [p for p in people if p not in completed_map]
    if incomplete_people:
        incomplete_townships[township] = incomplete_people

if incomplete_townships:
    for township, people in sorted(incomplete_townships.items()):
        people_str = '、'.join(people)
        report_lines.append(f'{township}：{people_str}。')
else:
    report_lines.append('无')

# 三违行为（只显示当天的）
report_lines.append('\n⚠️ 当天"三违"行为记录：')
all_sanwei = []
for name, records in sanwei_map.items():
    for record in records:
        all_sanwei.append(f'{name}：{record["details"]}')

if all_sanwei:
    report_lines.append(f'共发现 {len(all_sanwei)} 条"三违"行为：')
    for item in all_sanwei:
        report_lines.append(f'• {item}')
else:
    report_lines.append('无')

# 输出文字报告
text_report = '\n'.join(report_lines)
print(text_report)

# 保存文字报告
with open('统计报告_文字版.txt', 'w', encoding='utf-8') as f:
    f.write(text_report)

print('\n\n📁 文字报告已保存到：')
print('   [统计报告_文字版.txt](file:///f:/个人文件/workbuddy安全看板/三违管理/统计报告_文字版.txt)')

# ==================== 生成Excel表格版报告 ====================
wb = Workbook()
ws = wb.active
ws.title = f'{now.year}年{now.month}月巡检统计'

# 定义样式
header_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

content_font = Font(name='微软雅黑', size=10)
content_alignment = Alignment(horizontal='center', vertical='center')
completed_fill = PatternFill(start_color='C6E0B4', end_color='C6E0B4', fill_type='solid')
incomplete_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
penalty_fill = PatternFill(start_color='FF6B6B', end_color='FF6B6B', fill_type='solid')

# 生成自然月日期列表（月/日格式）
current_month = now.month
if current_month in [1, 3, 5, 7, 8, 10, 12]:
    days_in_month = 31
elif current_month in [4, 6, 9, 11]:
    days_in_month = 30
else:
    if (now.year % 4 == 0 and now.year % 100 != 0) or (now.year % 400 == 0):
        days_in_month = 29
    else:
        days_in_month = 28

date_list = []
for day in range(1, days_in_month + 1):
    date_list.append(f'{current_month}/{day}')

# 写入表头
headers = ['序号', '乡镇', '人员姓名']
headers.extend(date_list)
headers.append('总计')
headers.append('扣分')

for col_idx, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

# 写入人员数据
row_num = 2
seq = 1
all_personnel_sorted = sorted(all_personnel, key=lambda x: (x['township'], x['name']))
today_day = now.day

for person in all_personnel_sorted:
    ws.cell(row=row_num, column=1, value=seq).border = thin_border
    ws.cell(row=row_num, column=2, value=person['township']).border = thin_border
    ws.cell(row=row_num, column=3, value=person['name']).border = thin_border
    
    completed_count = 0
    incomplete_count = 0
    penalty_score = 0
    
    for col_idx, date_str in enumerate(date_list, start=4):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.border = thin_border
        cell.font = content_font
        cell.alignment = content_alignment
        
        name = person['name']
        is_completed = name in completed_map and date_str in completed_map[name]
        day_num = int(date_str.split('/')[1])
        is_past_or_today = day_num <= today_day
        
        if is_past_or_today:
            if is_completed:
                cell.value = '✓'
                cell.fill = completed_fill
                completed_count += 1
            else:
                cell.value = '❌'
                cell.fill = incomplete_fill
                incomplete_count += 1
                penalty_score += 1
    
    total_cell = ws.cell(row=row_num, column=4 + len(date_list))
    total_cell.value = f'✓{completed_count}天 ❌{incomplete_count}天'
    total_cell.font = Font(name='微软雅黑', size=10, bold=True)
    total_cell.alignment = content_alignment
    total_cell.border = thin_border
    
    penalty_cell = ws.cell(row=row_num, column=5 + len(date_list))
    penalty_cell.value = penalty_score
    penalty_cell.font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
    penalty_cell.alignment = content_alignment
    penalty_cell.border = thin_border
    if penalty_score > 0:
        penalty_cell.fill = penalty_fill
    
    for col_idx in range(1, 4):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.font = content_font
        cell.alignment = content_alignment
    
    row_num += 1
    seq += 1

# 添加统计汇总行
summary_row = row_num + 1
ws.cell(row=summary_row, column=1, value='').border = thin_border
ws.cell(row=summary_row, column=2, value='').border = thin_border

summary_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
summary_font = Font(name='微软雅黑', size=11, bold=True)

ws.cell(row=summary_row, column=3, value='汇总统计').border = thin_border
ws.cell(row=summary_row, column=3).font = summary_font
ws.cell(row=summary_row, column=3).alignment = content_alignment
ws.cell(row=summary_row, column=3).fill = summary_fill

total_completed = 0
total_incomplete = 0
total_penalty = 0

for col_idx in range(4, 4 + len(date_list)):
    completed_col = 0
    incomplete_col = 0
    
    for row_idx in range(2, row_num):
        cell_value = ws.cell(row=row_idx, column=col_idx).value
        if cell_value == '✓':
            completed_col += 1
            total_completed += 1
        elif cell_value == '❌':
            incomplete_col += 1
            total_incomplete += 1
    
    summary_cell = ws.cell(row=summary_row, column=col_idx)
    if completed_col > 0 or incomplete_col > 0:
        summary_cell.value = f'✓{completed_col}\n❌{incomplete_col}'
    else:
        summary_cell.value = ''
    summary_cell.font = Font(name='微软雅黑', size=9, bold=True)
    summary_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    summary_cell.border = thin_border
    summary_cell.fill = summary_fill

total_cell = ws.cell(row=summary_row, column=4 + len(date_list))
total_cell.value = f'✓共{total_completed}天 ❌共{total_incomplete}天'
total_cell.font = Font(name='微软雅黑', size=10, bold=True)
total_cell.alignment = content_alignment
total_cell.border = thin_border
total_cell.fill = summary_fill

for row_idx in range(2, row_num):
    penalty_value = ws.cell(row=row_idx, column=5 + len(date_list)).value
    if penalty_value and isinstance(penalty_value, int):
        total_penalty += penalty_value

penalty_summary_cell = ws.cell(row=summary_row, column=5 + len(date_list))
penalty_summary_cell.value = f'共扣{total_penalty}分'
penalty_summary_cell.font = Font(name='微软雅黑', size=10, bold=True)
penalty_summary_cell.alignment = content_alignment
penalty_summary_cell.border = thin_border
penalty_summary_cell.fill = summary_fill

# 调整列宽
ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 15
ws.column_dimensions['C'].width = 12
for col_idx in range(4, 4 + len(date_list)):
    ws.column_dimensions[get_column_letter(col_idx)].width = 8

ws.column_dimensions[get_column_letter(4 + len(date_list))].width = 16
ws.column_dimensions[get_column_letter(5 + len(date_list))].width = 10

ws.freeze_panes = 'D2'

# 保存Excel文件
excel_filename = f'统计报告_{now.year}年{now.month:02d}月.xlsx'
wb.save(excel_filename)

print(f'\n📊 Excel报告已保存到：')
print(f'   [{excel_filename}](file:///f:/个人文件/workbuddy安全看板/三违管理/{excel_filename})')