import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# 需要排除的人员名单
exclude_names = ['杨清华', '张和祥', '孙建栋', '谢俊章', '张万恒', '靳洪刚', '郭志强', '张卫新', '王臣祥', '李振恒', '张华平']

# 读取乡镇人员名单
personnel_df = pd.read_excel('乡镇人员名单（微信）.xlsx')

# 构建映射关系（只包含管理员，排除乡镇主管和指定人员）
name_to_wechat = {}
wechat_to_name = {}
township_to_people = {}
all_personnel = []

for idx, row in personnel_df.iterrows():
    duty = str(row.iloc[2]).strip()
    name = str(row.iloc[3]).strip()
    wechat = str(row.iloc[4]).strip()
    township = str(row.iloc[1]).strip()
    
    # 跳过乡镇主管
    if '乡镇主管' in duty:
        continue
    
    # 跳过指定要排除的人员
    if name in exclude_names:
        continue
    
    if name and name != 'nan':
        if wechat and wechat != 'nan':
            name_to_wechat[name] = wechat
            wechat_to_name[wechat] = name
        
        if township and township != 'nan':
            if township not in township_to_people:
                township_to_people[township] = []
            township_to_people[township].append(name)
            all_personnel.append({
                'township': township,
                'name': name,
                'wechat': wechat
            })

# 读取拍照巡检数据
inspection_df = pd.read_excel('拍照巡检 AI 识别.xlsx', header=None)

# 查找所有创建人（用于完成情况统计）
completed_map = {}  # {name: [dates]}

for idx in range(2, len(inspection_df)):
    creator = str(inspection_df.iloc[idx, 7]).strip()
    inspection_date = str(inspection_df.iloc[idx, 3]).strip()
    
    if creator and creator != 'nan':
        # 获取真实姓名
        name = creator
        if creator in wechat_to_name:
            name = wechat_to_name[creator]
        
        # 跳过被排除的人员
        if name in exclude_names:
            continue
        
        # 记录完成日期 - 转换为月/日格式
        if name not in completed_map:
            completed_map[name] = set()
        
        if inspection_date and inspection_date != 'nan':
            # 处理日期格式
            try:
                # 尝试解析 YYYY/MM/DD 格式
                if '/' in inspection_date:
                    date_obj = datetime.strptime(inspection_date, '%Y/%m/%d')
                    month_day = f'{date_obj.month}/{date_obj.day}'
                # 尝试解析 YYYY年MM月DD日 格式
                elif '年' in inspection_date and '月' in inspection_date and '日' in inspection_date:
                    year_part = inspection_date.split('年')[0]
                    month_part = inspection_date.split('年')[1].split('月')[0]
                    day_part = inspection_date.split('月')[1].split('日')[0]
                    month_day = f'{int(month_part)}/{int(day_part)}'
                else:
                    date_obj = datetime.strptime(inspection_date, '%Y-%m-%d')
                    month_day = f'{date_obj.month}/{date_obj.day}'
                
                completed_map[name].add(month_day)
            except:
                pass

# 获取当前日期和时间
now = datetime.now()
current_year = now.year
current_month = now.month

# 生成自然月日期列表（月/日格式）
date_list = []

# 计算当月天数
if current_month in [1, 3, 5, 7, 8, 10, 12]:
    days_in_month = 31
elif current_month in [4, 6, 9, 11]:
    days_in_month = 30
else:
    if (current_year % 4 == 0 and current_year % 100 != 0) or (current_year % 400 == 0):
        days_in_month = 29
    else:
        days_in_month = 28

for day in range(1, days_in_month + 1):
    date_list.append(f'{current_month}/{day}')

# ==================== 生成Excel表格版报告 ====================
wb = Workbook()
ws = wb.active
ws.title = f'{current_year}年{current_month}月巡检统计'

# 定义样式
header_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

content_font = Font(name='微软雅黑', size=10)
content_alignment = Alignment(horizontal='center', vertical='center')
completed_fill = PatternFill(start_color='C6E0B4', end_color='C6E0B4', fill_type='solid')
incomplete_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
penalty_fill = PatternFill(start_color='FF6B6B', end_color='FF6B6B', fill_type='solid')

# 写入表头
headers = ['序号', '乡镇', '人员姓名']
headers.extend(date_list)
headers.append('总计')
headers.append('扣分')

for col_idx, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

# 写入人员数据
row_num = 2
seq = 1
all_personnel_sorted = sorted(all_personnel, key=lambda x: (x['township'], x['name']))
today_day = now.day

for person in all_personnel_sorted:
    # 基本信息列
    ws.cell(row=row_num, column=1, value=seq).border = thin_border
    ws.cell(row=row_num, column=2, value=person['township']).border = thin_border
    ws.cell(row=row_num, column=3, value=person['name']).border = thin_border
    
    completed_count = 0
    incomplete_count = 0
    penalty_score = 0  # 扣分
    
    # 日期列
    for col_idx, date_str in enumerate(date_list, start=4):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.border = thin_border
        cell.font = content_font
        cell.alignment = content_alignment
        
        name = person['name']
        is_completed = name in completed_map and date_str in completed_map[name]
        
        # 解析日期中的天数
        day_num = int(date_str.split('/')[1])
        
        # 检查是否是今天之前的日期或今天
        is_past_or_today = day_num <= today_day
        
        if is_past_or_today:
            if is_completed:
                cell.value = '✓'
                cell.fill = completed_fill
                completed_count += 1
            else:
                cell.value = '❌'
                cell.fill = incomplete_fill
                incomplete_count += 1
                penalty_score += 1  # 每有一个❌扣1分
        # 未来的日期留空
    
    # 总计列
    total_cell = ws.cell(row=row_num, column=4 + len(date_list))
    total_cell.value = f'✓{completed_count}天 ❌{incomplete_count}天'
    total_cell.font = Font(name='微软雅黑', size=10, bold=True)
    total_cell.alignment = content_alignment
    total_cell.border = thin_border
    
    # 扣分列
    penalty_cell = ws.cell(row=row_num, column=5 + len(date_list))
    penalty_cell.value = penalty_score
    penalty_cell.font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
    penalty_cell.alignment = content_alignment
    penalty_cell.border = thin_border
    
    # 根据扣分设置颜色
    if penalty_score > 0:
        penalty_cell.fill = penalty_fill
    
    # 设置基本信息列的样式
    for col_idx in range(1, 4):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.font = content_font
        cell.alignment = content_alignment
    
    row_num += 1
    seq += 1

# 添加统计汇总行
summary_row = row_num + 1
ws.cell(row=summary_row, column=1, value='').border = thin_border
ws.cell(row=summary_row, column=2, value='').border = thin_border

summary_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
summary_font = Font(name='微软雅黑', size=11, bold=True)

ws.cell(row=summary_row, column=3, value='汇总统计').border = thin_border
ws.cell(row=summary_row, column=3).font = summary_font
ws.cell(row=summary_row, column=3).alignment = content_alignment
ws.cell(row=summary_row, column=3).fill = summary_fill

# 统计每列的完成和未完成数量
total_completed = 0
total_incomplete = 0
total_penalty = 0

for col_idx in range(4, 4 + len(date_list)):
    completed_col = 0
    incomplete_col = 0
    
    for row_idx in range(2, row_num):
        cell_value = ws.cell(row=row_idx, column=col_idx).value
        if cell_value == '✓':
            completed_col += 1
            total_completed += 1
        elif cell_value == '❌':
            incomplete_col += 1
            total_incomplete += 1
    
    summary_cell = ws.cell(row=summary_row, column=col_idx)
    if completed_col > 0 or incomplete_col > 0:
        summary_cell.value = f'✓{completed_col}\n❌{incomplete_col}'
    else:
        summary_cell.value = ''
    summary_cell.font = Font(name='微软雅黑', size=9, bold=True)
    summary_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    summary_cell.border = thin_border
    summary_cell.fill = summary_fill

# 总计列汇总
total_cell = ws.cell(row=summary_row, column=4 + len(date_list))
total_cell.value = f'✓共{total_completed}天 ❌共{total_incomplete}天'
total_cell.font = Font(name='微软雅黑', size=10, bold=True)
total_cell.alignment = content_alignment
total_cell.border = thin_border
total_cell.fill = summary_fill

# 扣分列汇总
for row_idx in range(2, row_num):
    penalty_value = ws.cell(row=row_idx, column=5 + len(date_list)).value
    if penalty_value and isinstance(penalty_value, int):
        total_penalty += penalty_value

penalty_summary_cell = ws.cell(row=summary_row, column=5 + len(date_list))
penalty_summary_cell.value = f'共扣{total_penalty}分'
penalty_summary_cell.font = Font(name='微软雅黑', size=10, bold=True)
penalty_summary_cell.alignment = content_alignment
penalty_summary_cell.border = thin_border
penalty_summary_cell.fill = summary_fill

# 调整列宽
ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 15
ws.column_dimensions['C'].width = 12
for col_idx in range(4, 4 + len(date_list)):
    ws.column_dimensions[get_column_letter(col_idx)].width = 8

# 设置最后两列宽一些
ws.column_dimensions[get_column_letter(4 + len(date_list))].width = 16
ws.column_dimensions[get_column_letter(5 + len(date_list))].width = 10

# 冻结窗格
ws.freeze_panes = 'D2'

# 保存Excel文件
excel_filename = f'统计报告_{current_year}年{current_month:02d}月.xlsx'
wb.save(excel_filename)

print(f'Excel报告已保存到：{excel_filename}')
print(f'\n✅ 已生成 {current_year}年{current_month}月 的巡检统计表')
print(f'📅 共 {len(date_list)} 天（日期格式：月/日）')
print(f'👥 共统计 {len(all_personnel_sorted)} 名管理员（已排除{len(exclude_names)}人）')
print(f'📊 汇总：已完成 {total_completed} 天次巡检，未完成 {total_incomplete} 天次巡检')
print(f'⚠️ 扣分统计：共扣 {total_penalty} 分（每有一个❌扣1分）')
print(f'\n已排除的人员：{", ".join(exclude_names)}')