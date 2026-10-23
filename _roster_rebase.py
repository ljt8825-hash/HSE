# -*- coding: utf-8 -*-
"""从现有（按花名册）结果文件重组：总网格长+一级网格长 → 公司管理（不按姓名，只合计）。
Sheet1/Sheet3(片区) 照搬；Sheet2/Sheet4 调整为「公司管理合计 + 二级/三级按姓名」。
数据全部从现有结果文件读取计算，不重新合并源文件。
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

src = r'G:\safety_vault\04_现场管理\三违\三位喂喂喂\新建文件夹(6)\2026年7月三违统计分析（按花名册）.xlsx'
out = src

# 读现有结果
s1 = pd.read_excel(src, sheet_name='1-片区巡查次数', header=1)
s1.columns = ['排名', '片区/班组', '网格长巡查次数']
s3 = pd.read_excel(src, sheet_name='3-片区三违数量', header=1)
s3.columns = ['排名', '片区/班组', '三违数量']

s2 = pd.read_excel(src, sheet_name='2-各等级网格长巡查按姓名', header=1)
s2.columns = ['网格长等级', '检查人姓名', '巡查次数']
s4 = pd.read_excel(src, sheet_name='4-各等级网格长三违按姓名', header=1)
s4.columns = ['网格长等级', '检查人姓名', '三违数量']

# 公司管理合计 = 片区表里"公司管理"行（本就=总网格长+一级网格长）
gongsi_inspect = int(s1.loc[s1['片区/班组'] == '公司管理', '网格长巡查次数'].iloc[0])
gongsi_viol = int(s3.loc[s3['片区/班组'] == '公司管理', '三违数量'].iloc[0])

# Sheet2 / Sheet4：公司管理(合计) + 二级 + 三级（保留原降序）
def rebase(df, valcol, gongsi):
    rows = [('公司管理', '（合计，不按姓名）', gongsi)]
    for lv in ['二级网格长', '三级网格长']:
        sub = df[df['网格长等级'] == lv]
        for _, r in sub.iterrows():
            rows.append((lv, r['检查人姓名'], int(r[valcol])))
    return rows

rows2 = rebase(s2, '巡查次数', gongsi_inspect)
rows4 = rebase(s4, '三违数量', gongsi_viol)

total_inspect = gongsi_inspect + int(s2[s2['网格长等级'].isin(['二级网格长', '三级网格长'])]['巡查次数'].sum())
total_violation = gongsi_viol + int(s4[s4['网格长等级'].isin(['二级网格长', '三级网格长'])]['三违数量'].sum())

dist2 = int(s2[s2['网格长等级']=='二级网格长']['巡查次数'].sum())
dist3 = int(s2[s2['网格长等级']=='三级网格长']['巡查次数'].sum())

print('公司管理 巡查:', gongsi_inspect, ' 三违:', gongsi_viol)
print('二级巡查:', dist2, ' 三级巡查:', dist3, ' 合计巡查:', total_inspect)
print('合计三违:', total_violation)

# ================= 写 Excel =================
wb = Workbook()
thin = Side(style='thin', color='CCCCCC')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill('solid', fgColor='DDEBF7')
title_font = Font(bold=True, size=14)
hdr_font = Font(bold=True, size=11)

def write_table(ws, title, headers, rows, col_widths=None):
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    c = ws.cell(1, 1); c.font = title_font; c.alignment = Alignment(horizontal='left')
    ws.append(headers)
    for i in range(1, len(headers)+1):
        cc = ws.cell(2, i); cc.font = hdr_font; cc.fill = hdr_fill
        cc.alignment = Alignment(horizontal='center', vertical='center'); cc.border = border
    for row in rows:
        ws.append(list(row))
    for r in range(3, 3 + len(rows)):
        for cc in range(1, len(headers)+1):
            ws.cell(r, cc).border = border
            ws.cell(r, cc).alignment = Alignment(horizontal='center' if cc != 2 else 'left')
    if col_widths:
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A3'

# 说明
ws0 = wb.active
ws0.title = '说明'
ws0.append(['2026年7月三违统计分析（依据网格长花名册重算）'])
ws0.append([])
ws0.append(['统计口径'])
ws0.append(['1. 巡查次数 = 有效记录行数（每行 = 一次网格长巡查/上报）；已剔除 2 行"检查人姓名"为空的占位行。'])
ws0.append(['2. 三违数量 = 仅统计「违章人姓名」非空的记录行数。'])
ws0.append(['3. 网格长等级一律以《网格长花名册》为权威口径，不再使用汇总表内错乱/空缺的"网格等级"列。'])
ws0.append(['4. 原表中"李中胜"为"李忠胜"笔误，已合并。'])
ws0.append(['5. 总网格长 + 一级网格长 合并归为「公司管理」，该类别不按姓名统计（只合计）。'])
ws0.append(['6. 按姓名统计仅含 二级网格长、三级网格长。'])
ws0.append([])
ws0.append(['总体汇总'])
ws0.append(['有效巡查记录数（剔除2行空姓名）', total_inspect])
ws0.append(['其中三违记录数（有违章人姓名）', total_violation])
ws0.append([])
ws0.append(['类别分布（花名册权威口径，全部有效行）'])
ws0.append(['  公司管理（总网格长+一级网格长）', gongsi_inspect])
ws0.append(['  二级网格长', dist2])
ws0.append(['  三级网格长', dist3])
ws0.append([])
ws0.append(['注记'])
ws0.append(['公司管理 = 总网格长 黄殿兴 + 一级网格长 刘金涛/张康/刘菲/韩晓亮/黄元同，不按姓名拆分。'])
ws0.append(['6 名花名册人员（丁新霞、孙小霞、孙海清、杨立卫、王晓燕、郑玉红，均为公园二级网格长）7月三违表中无巡查记录。'])
for i, w in enumerate([70, 60], 1):
    ws0.column_dimensions[get_column_letter(i)].width = w
ws0.cell(1, 1).font = Font(bold=True, size=15)

# Sheet1 片区巡查（照搬）
ws1 = wb.create_sheet('1-片区巡查次数')
write_table(ws1, '片区/班组 网格长巡查次数', ['排名', '片区/班组', '网格长巡查次数'],
            [(int(r['排名']), r['片区/班组'], int(r['网格长巡查次数'])) for _, r in s1.iterrows()],
            [8, 20, 18])

# Sheet2
ws2 = wb.create_sheet('2-巡查按姓名（公司管理合计）')
write_table(ws2, '网格长巡查次数（公司管理合计 + 二/三级按姓名）', ['网格长等级', '检查人姓名', '巡查次数'],
            rows2, [14, 20, 12])

# Sheet3 片区三违（照搬）
ws3 = wb.create_sheet('3-片区三违数量')
write_table(ws3, '片区/班组 三违数量（有违章人姓名）', ['排名', '片区/班组', '三违数量'],
            [(int(r['排名']), r['片区/班组'], int(r['三违数量'])) for _, r in s3.iterrows()],
            [8, 20, 18])

# Sheet4
ws4 = wb.create_sheet('4-三违按姓名（公司管理合计）')
write_table(ws4, '网格长三违数量（公司管理合计 + 二/三级按姓名）', ['网格长等级', '检查人姓名', '三违数量'],
            rows4, [14, 20, 12])

wb.save(out)
print('\n已写出:', out)
print('\nSheet2 前4行 + 行数:', len(rows2))
for r in rows2[:4]:
    print(r)
print('Sheet4 前4行 + 行数:', len(rows4))
for r in rows4[:4]:
    print(r)