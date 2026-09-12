import openpyxl, os, json
from collections import Counter

SRC = "Y:/月度考核/三违/9月/2026-09-1-6新版-按日期排序.xlsx"
OUT = "G:/hse-dashboard/server/scripts/_sync_0908_dbt/img_probe"
os.makedirs(OUT, exist_ok=True)

wb = openpyxl.load_workbook(SRC, read_only=False, data_only=True)
ws = wb["2026-09详细数据"]
imgs = ws._images
print("total images:", len(imgs))

info = []
for i, im in enumerate(imgs):
    anchor = im.anchor
    atype = type(anchor).__name__
    frm = getattr(anchor, "_from", None)
    row = getattr(frm, "row", None)
    col = getattr(frm, "col", None)
    # get bytes
    data = None
    try:
        data = im._data()
    except Exception:
        try:
            data = im._data
        except Exception:
            data = None
    if isinstance(data, bytes):
        head = data[:8]
        ext = "png" if head[:4] == b"\x89PNG" else ("jpg" if head[:3] == b"\xff\xd8\xff" else "bin")
        fn = f"{i:04d}_r{row}_c{col}_{atype}.{ext}"
        with open(os.path.join(OUT, fn), "wb") as f:
            f.write(data)
    else:
        fn = f"{i:04d}_r{row}_c{col}_{atype}.NODATA"
    # read 编号 / 姓名 / 检查日期 from that row (1-based)
    r1 = (row + 1) if isinstance(row, int) else None
    bianhao = ws.cell(row=r1, column=20).value if r1 else None   # 编号 col 20 (1-based)
    xingming = ws.cell(row=r1, column=4).value if r1 else None    # 姓名 col 4
    riqi = ws.cell(row=r1, column=7).value if r1 else None        # 日期 col 7
    info.append({"i": i, "anchor": atype, "row0": row, "col0": col, "file": fn,
                 "bytes": len(data) if isinstance(data, bytes) else 0,
                 "bianhao": bianhao, "xingming": xingming, "riqi": riqi})

print("by col0:", dict(Counter(x["col0"] for x in info)))
print("by anchor:", dict(Counter(x["anchor"] for x in info)))
print("with bianhao:", sum(1 for x in info if x["bianhao"]))
print("without bianhao:", sum(1 for x in info if not x["bianhao"]))
# sample
for x in info[:8]:
    print(x)

with open(os.path.join(OUT, "_map.json"), "w", encoding="utf-8") as f:
    json.dump(info, f, ensure_ascii=False, indent=1)
print("wrote _map.json, total", len(info))
