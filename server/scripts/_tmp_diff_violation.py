"""对比 Excel 真实 (name, vt) 与 store 中已存在的 (name, vt)，找到具体丢失/异常行。"""
import json
from openpyxl import load_workbook

def norm(v):
    if v is None: return ""
    if hasattr(v, 'strftime'): return v.strftime("%Y-%m-%d %H:%M:%S")
    return str(v).strip()

wb = load_workbook(r"G:/safety_vault/04_现场管理/2026交通专项行动/公司/驾驶员五年违章记录汇总.xlsx", data_only=True, read_only=True)
ws = wb.active
excel = []
for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    n = (r[2] or "").strip()
    vt = r[1]
    if not n: continue
    excel.append((i, n, norm(vt)))

# store rows
with open(r"G:/hse-dashboard/server/scripts/_store_dump.json", encoding="utf-8") as f:
    store = json.load(f)

# store 用 store.violationTime 已是 'YYYY-MM-DD HH:MM:SS' 与 norm 一致
store_keys = set((r['name'], r['violationTime'] or "") for r in store)
excel_keys_set = set()
excel_dup_idx = []
for tup in excel:
    k = (tup[1], tup[2])
    if k in excel_keys_set:
        excel_dup_idx.append(tup[0])
    else:
        excel_keys_set.add(k)

print(f"excel rows (non-empty): {len(excel)}")
print(f"excel unique (name, vt): {len(excel_keys_set)}")
print(f"excel internal duplicate rows: {len(excel_dup_idx)}")
print(f"store rows: {len(store)}, store unique (name, vt): {len(store_keys)}")

# store 中 (name, vt) 而不在 excel 的（手工/历史）
extra = sorted(store_keys - excel_keys_set)
print(f"extra in store (not in excel): {len(extra)}")
for e in extra[:50]:
    print(" ", e)

# excel 中 (name, vt) 而不在 store 的（应该写进去但没有）
missing = sorted(excel_keys_set - store_keys)
print(f"missing in store (excel says should be there): {len(missing)}")
for m in missing[:50]:
    print(" ", m)
