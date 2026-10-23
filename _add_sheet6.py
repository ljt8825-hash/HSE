# -*- coding: utf-8 -*-
"""给「2026年7月三违统计分析（按花名册）.xlsx」追加第6个sheet：
公司管理（总网格长+一级网格长）三违 按乡镇统计。
"""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

out = r'G:\safety_vault\04_现场管理\三违\三位喂喂喂\新建文件夹(6)\2026年7月三违统计分析（按花名册）.xlsx'

wb = load_workbook(out)
if '5-公司管理三违按乡镇' in wb.sheetnames:
    del wb['5-公司管理三违按乡镇']

ws = wb.create_sheet('5-公司管理三违按乡镇')

thin = Side(style='thin', color='CCCCCC')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill('solid', fgColor='DDEBF7')
title_font = Font(bold=True, size=14)
hdr_font = Font(bold=True, size=11)
center = Alignment(horizontal='center', vertical='center')
left = Alignment(horizontal='left', vertical='center', wrap_text=True)

headers = ['乡镇', '三违数量', '备注']

# 排序：乡镇按数量降序；车队（非乡镇）与未匹配放末尾
township_rows = [
    ('长官镇', 6, ''),
    ('大曹镇', 4, ''),
    ('相衙镇', 3, ''),
    ('保店镇', 2, ''),
    ('柴胡店镇', 1, ''),
    ('时集镇', 1, ''),
]
special_rows = [
    ('车队', 1, '非乡镇（杨小林，车队）'),
    ('未匹配', 2, '张书丽、杜凤兰（各花名册均未找到）'),
]
total = 20

# 标题
ws.append(['公司管理（总网格长 + 一级网格长）三违 按乡镇统计'])
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
c = ws.cell(1, 1); c.font = title_font; c.alignment = Alignment(horizontal='left')
ws.append([])

# 口径说明
note = ('口径：仅统计「违章人姓名」非空的公司管理记录，共 20 起；'
        '乡镇依据 7 月员工花名册「所属部门」归属（邢万俊=时集镇三班，按 7 月及前后花名册口径归「时集镇」）。')
ws.append([note])
ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=3)
ws.cell(3, 1).alignment = left
ws.cell(3, 1).font = Font(size=10, color='808080')
ws.append([])

# 表头
ws.append(headers)
hdr_row = ws.max_row
for i in range(1, len(headers) + 1):
    cc = ws.cell(hdr_row, i); cc.font = hdr_font; cc.fill = hdr_fill; cc.alignment = center; cc.border = border

# 数据行
for name, cnt, remark in township_rows + special_rows:
    ws.append([name, cnt, remark])

# 合计行
ws.append(['合计', total, ''])
total_row = ws.max_row

# 样式
data_start = hdr_row + 1
for r in range(data_start, total_row + 1):
    for cc in range(1, len(headers) + 1):
        cell = ws.cell(r, cc)
        cell.border = border
        cell.alignment = left if cc == 3 or (cc == 1) else center
    # 数量列居中
    ws.cell(r, 2).alignment = center
# 合计行加粗
for cc in range(1, len(headers) + 1):
    ws.cell(total_row, cc).font = Font(bold=True)

ws.column_dimensions[get_column_letter(1)].width = 18
ws.column_dimensions[get_column_letter(2)].width = 12
ws.column_dimensions[get_column_letter(3)].width = 40
ws.freeze_panes = 'A%d' % (data_start)

wb.save(out)
print('已写入 sheet「5-公司管理三违按乡镇」 ->', out)
print('sheets =', wb.sheetnames)