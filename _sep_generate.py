# -*- coding: utf-8 -*-
"""2026年9月 三违统计分析（依据网格长花名册，口径同7/8月）"""
import pandas as pd, re, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sep = r'c:\Users\honey\.trae-cn\attachments\6ab1177e5da3fb036f33ca96\ecfc3203-8e99-4173-ad3b-b31ea8df7651_c7099510-7049-47a2-990e-594c19dc21f5_2026年9月三违汇总（合并）.xlsx'
roster = r'c:\Users\honey\.trae-cn\attachments\6ab1177e5da3fb036f33ca96\ee67cd70-7f61-4017-879d-55c7d4008c7d_041a59d9-59c8-43fc-8874-6074005fa362_网格长花名册.xlsx'
out = r'G:\safety_vault\04_现场管理\三违\三位喂喂喂\新建文件夹(6)\2026年9月三违统计分析（按花名册）.xlsx'

md = pd.read_excel(sep, sheet_name='9月三违汇总')
md.columns = [str(c).strip() for c in md.columns]
md['检查人姓名'] = md['姓名'].astype(str).str.strip().replace({'李中胜': '李忠胜'})
md['分组名称'] = md['分组名称'].astype(str).str.strip().replace({'公司管理员': '公司管理'})

rb = pd.read_excel(roster); rb.columns = [str(c).strip() for c in rb.columns]
name2level = dict(zip(rb['姓名'].astype(str), rb['网格长等级']))
md['网格长等级'] = md['检查人姓名'].map(name2level)

def is_viol(x):
    s = str(x).strip()
    return s not in ('', '--', 'nan', 'None', '无', '——', '—')
md['是否三违'] = md['违章人姓名'].apply(is_viol)

d = md[md['网格长等级'].notna()].copy()
def classify(lv):
    return '公司管理' if lv in ('总网格长', '一级网格长') else lv
d['类别'] = d['网格长等级'].map(classify)

total_inspect = len(d)
total_viol = int(d['是否三违'].sum())
order = ['公司管理', '二级网格长', '三级网格长']

s1 = d.groupby('分组名称').size().sort_values(ascending=False).reset_index(name='网格长巡查次数')
s3 = d[d['是否三违']].groupby('分组名称').size().sort_values(ascending=False).reset_index(name='三违数量')

rows2, rows4 = [], []
for cat in order:
    sub = d[d['类别'] == cat]; subv = sub[sub['是否三违']]
    if cat == '公司管理':
        rows2.append((cat, '（合计，不按姓名）', int(len(sub))))
        rows4.append((cat, '（合计，不按姓名）', int(len(subv))))
    else:
        for n, c in sub.groupby('检查人姓名').size().sort_values(ascending=False).items():
            rows2.append((cat, n, int(c)))
        for n, c in subv.groupby('检查人姓名').size().sort_values(ascending=False).items():
            rows4.append((cat, n, int(c)))

# 乡镇映射
def town(dept):
    dd = str(dept).strip()
    return re.sub(r'(片区|一班|二班|三班|四班|五班|六班|七班|八班|九班|十班|三四班|六七班|一二班|二三班)$', '', dd)

emp_frames = []
base = r'G:\safety_vault\08_保险及人员管理\花名册'
for f, sheets in [(os.path.join(base, '员工信息维护_20260910140825在职版.xlsx'), ['员工信息维护']),
                  (os.path.join(base, '2026.9.1.xlsx'), ['员工信息维护'])]:
    for sh in sheets:
        try:
            df = pd.read_excel(f, sheet_name=sh)
            df.columns = [str(c).strip() for c in df.columns]
            if '姓名' in df.columns and '所属部门' in df.columns:
                emp_frames.append(df[['姓名', '所属部门']])
        except Exception:
            pass
emp = pd.concat(emp_frames).dropna(subset=['姓名']).drop_duplicates(subset=['姓名'])
emp['姓名'] = emp['姓名'].astype(str).str.strip()
name2town = dict(zip(emp['姓名'], emp['所属部门'].map(town)))

def clean_viol(s):
    return re.sub(r'[（(].*?[)）]', '', str(s).strip()).strip()
typo_map = {'姚奎岭': '姚奎领', '李凤丽': '李风丽', '李建方': '李建芳', '范志新': '范治新',
            '李兰凤': '李兰风', '杜瑞兰': '杜润兰'}

gv = d[(d['类别'] == '公司管理') & d['是否三违']].copy()
gv['违章人'] = gv['违章人姓名'].map(clean_viol).replace(typo_map)
gv['乡镇'] = gv['违章人'].map(name2town).fillna('未匹配')

town_df = gv.groupby('乡镇').size().reset_index(name='三违数量').sort_values('三违数量', ascending=False).reset_index(drop=True)
town_rows = [(r['乡镇'], int(r['三违数量'])) for _, r in town_df.iterrows()]
town_total = int(len(gv))

zg = d[d['网格长等级'] == '总网格长']; yj = d[d['网格长等级'] == '一级网格长']
zg_n = sorted(zg['检查人姓名'].unique()); yj_n = sorted(yj['检查人姓名'].unique())

print('行:', len(md), ' 有效:', total_inspect, ' 三违:', total_viol)
print('类别:', dict(sorted(d['类别'].value_counts().items())))
print('公司管理=总(%s)%d + 一级(%s)%d = %d 三违%d' % ('、'.join(zg_n), len(zg), '、'.join(yj_n), len(yj), len(zg)+len(yj), int(d[d['类别']=='公司管理']['是否三违'].sum())))
print('乡镇:', dict(town_rows), '合计', town_total)

# ============ 写 Excel ============
wb = Workbook()
thin = Side(style='thin', color='CCCCCC'); border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill('solid', fgColor='DDEBF7')
title_font = Font(bold=True, size=14); hdr_font = Font(bold=True, size=11)
center = Alignment(horizontal='center', vertical='center')

def write_table(ws, title, headers, rows, col_widths=None):
    ws.append([title]); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1).font = title_font; ws.cell(1, 1).alignment = Alignment(horizontal='left')
    ws.append(headers)
    for i in range(1, len(headers)+1):
        cc = ws.cell(2, i); cc.font = hdr_font; cc.fill = hdr_fill; cc.alignment = center; cc.border = border
    for row in rows:
        ws.append(list(row))
    for r in range(3, 3 + len(rows)):
        for cc in range(1, len(headers)+1):
            ws.cell(r, cc).border = border
            ws.cell(r, cc).alignment = Alignment(horizontal='center' if cc != 2 else 'left')
    if col_widths:
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A3'

ws0 = wb.active; ws0.title = '说明'
ws0.append(['2026年9月三违统计分析（依据网格长花名册重算）'])
ws0.append([]); ws0.append(['统计口径'])
for t in [
    '1. 数据覆盖 2026-09-01 ~ 09-18（18 天，截至 9/21 未满月）。',
    '2. 巡查次数 = 有效记录行数（每行 = 一次网格长巡查上报）；本表无空行。',
    '3. 三违数量 = 仅统计「违章人姓名」非空/非“--”的记录行数。',
    '4. 网格长等级以《网格长花名册》权威口径（同7/8月），台账「网格等级」列与花名册 94 行不一致。',
    '5. 原表「李中胜」为「李忠胜」笔误已合并（若出现）。',
    '6. 总网格长 + 一级网格长 合并「公司管理」不按姓名；按姓名仅 二级/三级。',
    '7. 分组名称中「公司管理员」与「公司管理」为同一批人跨批次标法不一致，已归一为「公司管理」。',
]:
    ws0.append([t])
ws0.append([]); ws0.append(['总体汇总'])
ws0.append(['有效巡查记录数', total_inspect])
ws0.append(['其中三违记录数（有违章人姓名）', total_viol])
ws0.append([]); ws0.append(['类别分布（花名册权威口径）'])
for cat in order:
    ws0.append(['  ' + cat, int(d['类别'].value_counts().get(cat, 0))])
ws0.append([]); ws0.append(['注记'])
ws0.append(['公司管理 = 总网格长 %s（%d 行）+ 一级网格长 %s（%d 行），合计 %d 行、三违 %d 起。'
            % ('、'.join(zg_n), len(zg), '、'.join(yj_n), len(yj), len(zg)+len(yj),
               int(d[d['类别']=='公司管理']['是否三违'].sum()))])
ws0.append(['公司管理三违按乡镇：违章人姓名做了清洗与笔误归并（范志新→范治新、李兰凤→李兰风、杜瑞兰→杜润兰、王春灵去备注等），其余直接按 9 月员工花名册「所属部门」归乡镇。'])
for i, w in enumerate([70, 60], 1):
    ws0.column_dimensions[get_column_letter(i)].width = w
ws0.cell(1, 1).font = Font(bold=True, size=15)

ws1 = wb.create_sheet('1-片区巡查次数')
write_table(ws1, '片区/班组 网格长巡查次数', ['排名', '片区/班组', '网格长巡查次数'],
            [(i+1, r['分组名称'], int(r['网格长巡查次数'])) for i, r in s1.iterrows()], [8, 20, 18])

ws2 = wb.create_sheet('2-巡查按姓名（公司管理合计）')
write_table(ws2, '网格长巡查次数（公司管理合计 + 二/三级按姓名）', ['网格长等级', '检查人姓名', '巡查次数'],
            rows2, [14, 20, 12])

ws3 = wb.create_sheet('3-片区三违数量')
write_table(ws3, '片区/班组 三违数量（有违章人姓名）', ['排名', '片区/班组', '三违数量'],
            [(i+1, r['分组名称'], int(r['三违数量'])) for i, r in s3.iterrows()], [8, 20, 18])

ws4 = wb.create_sheet('4-三违按姓名（公司管理合计）')
write_table(ws4, '网格长三违数量（公司管理合计 + 二/三级按姓名）', ['网格长等级', '检查人姓名', '三违数量'],
            rows4, [14, 20, 12])

ws5 = wb.create_sheet('5-公司管理三违按乡镇')
write_table(ws5, '公司管理（总网格长 + 一级网格长）三违 按乡镇统计', ['乡镇', '三违数量'],
            town_rows + [('合计', town_total)], [18, 12])

wb.save(out)
print('\n已写出:', out)
print('\nSheet5 乡镇:'); [print(r) for r in town_rows]