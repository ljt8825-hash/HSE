# -*- coding: utf-8 -*-
import openpyxl, os
roster = r'c:\Users\honey\.trae-cn\attachments\6ab1177e5da3fb036f33ca96\ee67cd70-7f61-4017-879d-55c7d4008c7d_041a59d9-59c8-43fc-8874-6074005fa362_网格长花名册.xlsx'
srcdir = r'G:\safety_vault\04_现场管理\三违\三位喂喂喂\新建文件夹(6)'

wb = openpyxl.load_workbook(roster, data_only=True)
ws = wb['Sheet1']
names = [(ws.cell(r,2).value, ws.cell(r,6).value) for r in range(2, ws.max_row+1)]
print('花名册含 忠/中 相关姓名:')
for n, lv in names:
    if n and ('忠' in str(n) or '中' in str(n)):
        print('   ', repr(n), '->', repr(lv))
wb.close()

# 7.28 源文件 网格等级列
f = os.path.join(srcdir, '每日“三违”检查汇总2026.7.28.xlsx')
wb = openpyxl.load_workbook(f, data_only=True)
ws = wb['详细数据统计（按预设名单顺序排序）']
hdr = [ws.cell(2,c).value for c in range(1, ws.max_column+1)]
print('\n7.28 表头:', hdr)
# 数据前三行 前7列
for r in range(3, 6):
    vals = [ws.cell(r,c).value for c in range(1, 8)]
    print(f'  数据R{r}:', vals)
wb.close()