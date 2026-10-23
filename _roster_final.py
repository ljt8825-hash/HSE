# -*- coding: utf-8 -*-
"""依据网格长花名册重算7月三违统计。
口径：总网格长 + 一级网格长 合并为「公司管理」（不按姓名，只合计）；二级/三级按姓名。
李中胜→李忠胜合并；剔除空检查人姓名占位行。
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

roster = r'c:\Users\honey\.trae-cn\attachments\6ab1177e5da3fb036f33ca96\ee67cd70-7f61-4017-879d-55c7d4008c7d_041a59d9-59c8-43fc-8874-6074005fa362_网格长花名册.xlsx'
merged = r'G:\safety_vault\04_现场管理\三违\三位喂喂喂\新建文件夹(6)\2026年7月三违汇总（合并）.xlsx'
out = r'G:\safety_vault\04_现场管理\三违\三位喂喂喂\新建文件夹(6)\2026年7月三违统计分析（按花名册）.xlsx'

rb = pd.read_excel(roster)
rb.columns = [str(c).strip() for c in rb.columns]
md = pd.read_excel(merged)
md.columns = [str(c).strip() for c in md.columns]

name2level = dict(zip(rb['姓名'], rb['网格长等级']))
md['检查人姓名'] = md['检查人姓名'].replace({'李中胜': '李忠胜'})

d = md[md['检查人姓名'].notna() & (md['检查人姓名'] != '')].copy()
d['网格长等级'] = d['检查人姓名'].map(name2level)
d['是否三违'] = d['违章人姓名'].notna() & (d['违章人姓名'].astype(str).str.strip() != '')

# 归类：总网格长 + 一级网格长 -> 公司管理
def classify(lv):
    return '公司管理' if lv in ('总网格长', '一级网格长') else lv
d['类别'] = d['网格长等级'].map(classify)

total_inspect = len(d)
total_violation = int(d['是否三违'].sum())

order = ['公司管理', '二级网格长', '三级网格长']

# Sheet1 片区巡查（片区/班组字段，公司管理本就=总+一级）
s1 = d.groupby('片区/班组').size().sort_values(ascending=False).reset_index(name='网格长巡查次数')
# Sheet3 片区三违
s3 = d[d['是否三违']].groupby('片区/班组').size().sort_values(ascending=False).reset_index(name='三违数量')

# Sheet2 按姓名巡查：公司管理=合计一行；二级/三级按姓名
rows2 = []
for cat in order:
    sub = d[d['类别'] == cat]
    if cat == '公司管理':
        rows2.append((cat, '（合计，不按姓名）', int(len(sub))))
    else:
        g = sub.groupby('检查人姓名').size().sort_values(ascending=False)
        for n, c in g.items():
            rows2.append((cat, n, int(c)))

# Sheet4 按姓名三违：公司管理=合计一行；二级/三级按姓名
rows4 = []
for cat in order:
    sub = d[(d['类别'] == cat) & d['是否三违']]
    if cat == '公司管理':
        rows4.append((cat, '（合计，不按姓名）', int(len(sub))))
    else:
        g = sub.groupby('检查人姓名').size().sort_values(ascending=False)
        for n, c in g.items():
            rows4.append((cat, n, int(c)))

dist = d['类别'].value_counts()
zg = d[d['网格长等级'] == '总网格长']
yj = d[d['网格长等级'] == '一级网格长']
zg_n = zg['检查人姓名'].unique().tolist()
yj_n = sorted(yj['检查人姓名'].unique().tolist())

print('有效巡查记录:', total_inspect, ' 三违:', total_violation)
print('类别分布:', dict(sorted(dist.items(), key=lambda x: order.index(x[0]))))
print('公司管理 = 总网格长(%s)%d 行 + 一级网格长(%s)%d 行 = %d 行' % (
    '、'.join(zg_n), len(zg), '、'.join(yj_n), len(yj), len(zg)+len(yj)))
print('公司管理三违:', int(d[d['类别']=='公司管理']['是否三违'].sum()))

# ================= 写 Excel =================
wb = Workbook()
thin = Side(style='thin', color='CCCCCC')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill('solid', fgColor='DDEBF7')
title_font = Font(bold=True, size=14)
hdr_font = Font(bold=True, size=11)
center = Alignment(horizontal='center', vertical='center')

def write_table(ws, title, headers, rows, col_widths=None):
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    c = ws.cell(1, 1); c.font = title_font; c.alignment = Alignment(horizontal='left')
    ws.append(headers)
    for i in range(1, len(headers)+1):
        cc = ws.cell(2, i); cc.font = hdr_font; cc.fill = hdr_fill; cc.alignment = center; cc.border = border
    for row in rows:
        ws.append(list(row))
    for r in range(3, 3 + len(rows)):
        for cc in range(1, len(headers)+1):
            ws.cell(r, cc).border = border
            ws.cell(r, cc).alignment = Alignment(horizontal='center' if cc != 2 else 'left')
    if col_widths:
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A3'

# 说明
ws0 = wb.active
ws0.title = '说明'
ws0.append(['2026年7月三违统计分析（依据网格长花名册重算）'])
ws0.append([])
ws0.append(['统计口径'])
ws0.append(['1. 巡查次数 = 有效记录行数（每行 = 一次网格长巡查/上报）；已剔除 2 行"检查人姓名"为空的占位行。'])
ws0.append(['2. 三违数量 = 仅统计「违章人姓名」非空的记录行数。'])
ws0.append(['3. 网格长等级一律以《网格长花名册》为权威口径，不再使用汇总表内错乱/空缺的"网格等级"列。'])
ws0.append(['4. 原表中"李中胜"为"李忠胜"笔误，已合并。'])
ws0.append(['5. 总网格长 + 一级网格长 合并归为「公司管理」，该类别不按姓名统计（只合计）。'])
ws0.append(['6. 按姓名统计仅含 二级网格长、三级网格长。'])
ws0.append([])
ws0.append(['总体汇总'])
ws0.append(['有效巡查记录数（剔除2行空姓名）', total_inspect])
ws0.append(['其中三违记录数（有违章人姓名）', total_violation])
ws0.append([])
ws0.append(['类别分布（花名册权威口径，全部有效行）'])
for cat in order:
    ws0.append(['  ' + cat, int(dist.get(cat, 0))])
ws0.append([])
ws0.append(['注记'])
ws0.append(['公司管理 = 总网格长 %s（%d 行） + 一级网格长 %s（%d 行），合计 %d 行、三违 %d 起；不按姓名拆分。'
            % ('、'.join(zg_n), len(zg), '、'.join(yj_n), len(yj), len(zg)+len(yj),
               int(d[d['类别']=='公司管理']['是否三违'].sum()))])
ws0.append(['6 名花名册人员（丁新霞、孙小霞、孙海清、杨立卫、王晓燕、郑玉红，均为公园二级网格长）7月三违表中无巡查记录。'])
for i, w in enumerate([70, 60], 1):
    ws0.column_dimensions[get_column_letter(i)].width = w
ws0.cell(1, 1).font = Font(bold=True, size=15)

# Sheet1
ws1 = wb.create_sheet('1-片区巡查次数')
write_table(ws1, '片区/班组 网格长巡查次数', ['排名', '片区/班组', '网格长巡查次数'],
            [(i+1, r['片区/班组'], int(r['网格长巡查次数'])) for i, r in s1.iterrows()],
            [8, 20, 18])

# Sheet2
ws2 = wb.create_sheet('2-巡查按姓名（公司管理合计）')
write_table(ws2, '网格长巡查次数（公司管理合计 + 二/三级按姓名）', ['网格长等级', '检查人姓名', '巡查次数'],
            rows2, [14, 20, 12])

# Sheet3
ws3 = wb.create_sheet('3-片区三违数量')
write_table(ws3, '片区/班组 三违数量（有违章人姓名）', ['排名', '片区/班组', '三违数量'],
            [(i+1, r['片区/班组'], int(r['三违数量'])) for i, r in s3.iterrows()],
            [8, 20, 18])

# Sheet4
ws4 = wb.create_sheet('4-三违按姓名（公司管理合计）')
write_table(ws4, '网格长三违数量（公司管理合计 + 二/三级按姓名）', ['网格长等级', '检查人姓名', '三违数量'],
            rows4, [14, 20, 12])

wb.save(out)
print('\n已写出:', out)
print('\nSheet2 明细:')
for r in rows2:
    print(r)
print('\nSheet4 明细:')
for r in rows4:
    print(r)