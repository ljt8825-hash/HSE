# -*- coding: utf-8 -*-
import openpyxl, os
from collections import Counter

roster = r'c:\Users\honey\.trae-cn\attachments\6ab1177e5da3fb036f33ca96\ee67cd70-7f61-4017-879d-55c7d4008c7d_041a59d9-59c8-43fc-8874-6074005fa362_网格长花名册.xlsx'
data = r'c:\Users\honey\.trae-cn\attachments\6ab1177e5da3fb036f33ca96\59495499-c858-45fd-86f0-a1eda81750a2_e451bd0c-4f00-405d-ad41-14c545831138_2026年7月三违汇总（合并）.xlsx'

wb = openpyxl.load_workbook(roster, data_only=True)
print('=== 网格长花名册 ===')
for ws in wb.worksheets:
    print(f'  sheet={ws.title!r} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}')
    print('  merged:', list(ws.merged_cells.ranges)[:10])
    # 前6行
    for r in range(1, min(ws.max_row, 6)+1):
        vals=[f'{openpyxl.utils.get_column_letter(c)}={ws.cell(r,c).value!r}' for c in range(1, ws.max_column+1) if ws.cell(r,c).value is not None]
        print(f'   R{r}:', ' | '.join(vals))
wb.close()

print()
wb = openpyxl.load_workbook(data, data_only=True)
ws = wb.active
print('=== 三违汇总表表头 ===')
print([ws.cell(1,c).value for c in range(1, ws.max_column+1)])
wb.close()