import openpyxl
f = r'G:\safety_vault\04_现场管理\三违\三位喂喂喂\新建文件夹(6)\2026年7月三违统计分析.xlsx'
wb = openpyxl.load_workbook(f, data_only=True)
for ws in wb.worksheets:
    print('=== Sheet:', ws.title, ' dims:', ws.dimensions, ' max_row:', ws.max_row, ' max_col:', ws.max_column, '===')
    for r in range(1, min(ws.max_row, 90)+1):
        vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column, 8)+1)]
        if any(v is not None for v in vals):
            print(r, vals)
wb.close()